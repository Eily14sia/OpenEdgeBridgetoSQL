"""
Generate styled PDFs (and optionally DOCX) from the docs markdown files.

Usage:
    python docs/generate_client_pdf.py

Outputs:
    docs/OpenEdgeBridge_Client_Setup_Guide.pdf
    docs/OpenEdgeBridge_Client_Setup_Guide.docx
    docs/OpenEdgeBridge_Documentation.pdf
"""
import os
import re

import markdown
from xhtml2pdf import pisa

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MD_FILES = [
    ("OpenEdgeBridge_Client_Setup_Guide.md",
     "OpenEdgeBridge_Client_Setup_Guide.pdf",
     "OpenEdgeBridge_Client_Setup_Guide.docx"),
    ("OpenEdgeBridge_Documentation.md",
     "OpenEdgeBridge_Documentation.pdf",
     None),
]

CSS = """
@page {
    size: letter;
    margin: 1.8cm 1.6cm 2cm 1.6cm;
    @frame footer {
        -pdf-frame-content: footer;
        bottom: 0.7cm;
        height: 1cm;
        text-align: center;
        font-size: 8pt;
        color: #888;
    }
}

body {
    font-family: Helvetica, Arial, sans-serif;
    font-size: 9.5pt;
    line-height: 1.45;
    color: #1f2937;
}

h1 {
    font-size: 19pt;
    color: #0f3d6e;
    border-bottom: 2.5pt solid #0f3d6e;
    padding-bottom: 6px;
    margin-bottom: 10px;
}

h2 {
    font-size: 13.5pt;
    color: #0f3d6e;
    background-color: #eef4fb;
    padding: 4px 8px;
    border-left: 4pt solid #0f3d6e;
    margin-top: 18px;
    page-break-after: avoid;
}

h3 {
    font-size: 11pt;
    color: #1668b0;
    margin-top: 12px;
    page-break-after: avoid;
}

h4 {
    font-size: 10pt;
    color: #333;
    margin-top: 10px;
}

p { margin: 6px 0; }

ul, ol { margin: 6px 0 6px 18px; padding: 0; }
li { margin: 3px 0; }

strong { color: #0f3d6e; }

pre {
    font-family: "Courier", monospace;
    font-size: 8pt;
    background-color: #f4f6f8;
    border: 0.6pt solid #d1d5db;
    border-left: 3pt solid #0f3d6e;
    padding: 8px 10px;
    margin: 8px 0;
    white-space: pre-wrap;
    page-break-inside: avoid;
}

code {
    font-family: "Courier", monospace;
    font-size: 8.5pt;
    background-color: #eef1f4;
    padding: 0 2px;
}

pre code { background-color: transparent; padding: 0; }

table {
    border-collapse: collapse;
    width: 100%;
    margin: 10px 0;
    font-size: 8.5pt;
    page-break-inside: auto;
}

th {
    background-color: #0f3d6e;
    color: #ffffff;
    font-weight: bold;
    padding: 5px 7px;
    border: 0.6pt solid #0f3d6e;
    text-align: left;
}

td {
    padding: 4px 7px;
    border: 0.6pt solid #cbd5e1;
    vertical-align: top;
}

tr:nth-child(even) td { background-color: #f2f6fb; }

blockquote {
    background-color: #fff8e6;
    border-left: 4pt solid #eab308;
    padding: 6px 10px;
    margin: 8px 0;
    font-size: 9pt;
}

hr {
    border: none;
    border-top: 1pt solid #d1d5db;
    margin: 14px 0;
}

#footer {
    font-size: 8pt;
    color: #888888;
    text-align: center;
}
"""


def build_html(md_text: str, footer_title: str) -> str:
    body = markdown.markdown(
        md_text,
        extensions=["tables", "fenced_code", "nl2br"],
        output_format="html5",
    )
    return f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8"><style>{CSS}</style></head>
<body>
{body}
<div id="footer">OpenEdge Bridge — {footer_title} &middot; Page <pdf:pagenumber /> of <pdf:pagecount /></div>
</body>
</html>"""


def make_pdf(md_path: str, pdf_path: str, footer_title: str) -> None:
    with open(md_path, encoding="utf-8") as f:
        md_text = f.read()
    html = build_html(md_text, footer_title)
    with open(pdf_path, "wb") as out:
        pisa.CreatePDF(html, dest=out, encoding="utf-8")
    print(f"PDF created: {pdf_path}")


def make_docx(md_path: str, docx_path: str) -> None:
    """Simple DOCX conversion: HTML -> DOCX via python-docx (if available)."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from html.parser import HTMLParser
    except ImportError:
        print("python-docx not installed; skipping DOCX generation.")
        return

    with open(md_path, encoding="utf-8") as f:
        md_text = f.read()

    # Convert markdown to HTML, then strip tags into structured text lines.
    html = markdown.markdown(
        md_text,
        extensions=["tables", "fenced_code", "nl2br"],
        output_format="html5",
    )

    class BlockParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.blocks = []      # list of (kind, text)
            self.cur = []         # current inline text pieces
            self.cur_kind = "p"
            self.in_code = False
            self.in_table = False
            self.table_rows = []
            self.cur_row = []
            self.cur_cell = []
            self.in_cell = False

        def _flush(self):
            text = "".join(self.cur).strip()
            if text:
                if self.cur_kind == "code":
                    self.blocks.append(("code", text))
                elif self.cur_kind == "th" or self.cur_kind == "td":
                    self.cur_cell.append("".join(self.cur).strip())
                else:
                    self.blocks.append((self.cur_kind, text))
            self.cur = []

        def handle_starttag(self, tag, attrs):
            if tag in ("pre",):
                self._flush()
                self.cur_kind = "code"
                self.in_code = True
            elif tag == "table":
                self.in_table = True
                self.table_rows = []
            elif tag == "tr":
                self.cur_row = []
            elif tag in ("th", "td"):
                self._flush()
                self.cur_cell = []
                self.in_cell = True
                self.cur_kind = tag
            elif tag in ("h1", "h2", "h3", "h4"):
                self._flush()
                self.cur_kind = tag
            elif tag in ("p", "li", "blockquote"):
                self._flush()
                self.cur_kind = tag
            elif tag == "br":
                self.cur.append("\n")
            elif tag in ("strong", "b"):
                self.cur.append("")
            elif tag == "code" and not self.in_code:
                self.cur.append("`")

        def handle_endtag(self, tag):
            if tag == "pre":
                self._flush()
                self.cur_kind = "p"
                self.in_code = False
            elif tag in ("th", "td"):
                self.cur_row.append("".join(self.cur_cell).strip())
                self.cur_cell = []
                self.in_cell = False
                self.cur_kind = "p"
            elif tag == "tr":
                self.table_rows.append(self.cur_row)
            elif tag == "table":
                self.blocks.append(("table", self.table_rows))
                self.in_table = False
            elif tag in ("h1", "h2", "h3", "h4", "p", "li", "blockquote"):
                self._flush()
                self.cur_kind = "p"
            elif tag == "code" and not self.in_code:
                self.cur.append("`")

        def handle_data(self, data):
            if data.strip() or self.cur_kind == "code":
                self.cur.append(data)

    parser = BlockParser()
    parser.feed(html)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    def add_heading(text, level):
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x0F, 0x3D, 0x6E)
        return h

    for kind, content in parser.blocks:
        if kind == "h1":
            add_heading(content, 1)
        elif kind == "h2":
            add_heading(content, 2)
        elif kind == "h3":
            add_heading(content, 3)
        elif kind == "h4":
            add_heading(content, 4)
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            p.paragraph_format.left_indent = Inches(0.3)
        elif kind == "table":
            rows = content
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for i, row in enumerate(rows):
                    for j, cell_text in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = cell_text
            doc.add_paragraph()
        else:  # p, li, blockquote
            p = doc.add_paragraph(content)
            if kind == "blockquote":
                p.paragraph_format.left_indent = Inches(0.3)
            elif kind == "li":
                p.style = "List Bullet"

    doc.save(docx_path)
    print(f"DOCX created: {docx_path}")


if __name__ == "__main__":
    for md_name, pdf_name, docx_name in MD_FILES:
        md_path = os.path.join(BASE_DIR, md_name)
        pdf_path = os.path.join(BASE_DIR, pdf_name)

        footer_title = "Client-Side Setup Guide" if "Setup_Guide" in md_name else "Documentation"
        make_pdf(md_path, pdf_path, footer_title)

        if docx_name:
            make_docx(md_path, os.path.join(BASE_DIR, docx_name))
